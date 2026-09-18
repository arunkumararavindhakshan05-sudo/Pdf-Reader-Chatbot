"""Turn any supported document into page-aware, searchable chunks.

Every loader returns the same shape as the original PDF pipeline,
``(chunks, section_count)``, so the retriever and RAG service work unchanged
for Word, Excel, CSV, PowerPoint and plain-text files.
"""

import csv
import io
from collections.abc import Callable
from pathlib import PurePath

from langchain_text_splitters import RecursiveCharacterTextSplitter

from src.document_processor import DocumentChunk, extract_chunks

LoaderResult = tuple[list[DocumentChunk], int]

MAX_ROWS_PER_CHUNK = 40
PARAGRAPHS_PER_SECTION = 25


class UnsupportedFileTypeError(ValueError):
    """Raised when a file extension has no registered loader."""


def _make_splitter(
    chunk_size: int, chunk_overlap: int
) -> RecursiveCharacterTextSplitter:
    if chunk_size <= 0:
        raise ValueError("Chunk size must be greater than zero.")

    if chunk_overlap < 0 or chunk_overlap >= chunk_size:
        raise ValueError(
            "Chunk overlap must be non-negative and smaller than chunk size."
        )

    return RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )


def _split_sections(
    sections: list[tuple[str, str]],
    chunk_size: int,
    chunk_overlap: int,
) -> list[DocumentChunk]:
    """Split ``(location, text)`` sections into numbered chunks."""
    splitter = _make_splitter(chunk_size, chunk_overlap)
    chunks: list[DocumentChunk] = []

    for section_number, (location, text) in enumerate(sections, start=1):
        cleaned = text.strip()

        if not cleaned:
            continue

        for chunk_number, chunk_text in enumerate(
            splitter.split_text(cleaned),
            start=1,
        ):
            chunks.append(
                {
                    "text": chunk_text,
                    "page": section_number,
                    "chunk": chunk_number,
                    "location": location,
                }
            )

    return chunks


def _format_cell(value: object) -> str:
    if value is None:
        return ""

    return str(value).strip()


def _rows_to_sections(
    label: str,
    rows: list[list[str]],
) -> list[tuple[str, str]]:
    """Group table rows into sections that repeat the header row.

    Repeating the header in every section keeps each chunk self-explanatory,
    so a retrieved block of rows still says which column is which.
    """
    rows = [row for row in rows if any(cell for cell in row)]

    if not rows:
        return []

    header, body = rows[0], rows[1:]
    header_line = " | ".join(header)

    if not body:
        return [(f"{label}, row 1", header_line)]

    sections: list[tuple[str, str]] = []

    for start in range(0, len(body), MAX_ROWS_PER_CHUNK):
        block = body[start : start + MAX_ROWS_PER_CHUNK]
        first_row = start + 2
        last_row = start + 1 + len(block)
        lines = [header_line, *(" | ".join(row) for row in block)]
        sections.append((f"{label}, rows {first_row}-{last_row}", "\n".join(lines)))

    return sections


def load_pdf(data: bytes, chunk_size: int, chunk_overlap: int) -> LoaderResult:
    chunks, page_count = extract_chunks(data, chunk_size, chunk_overlap)

    for chunk in chunks:
        chunk["location"] = f"Page {chunk['page']}"

    return chunks, page_count


def load_docx(data: bytes, chunk_size: int, chunk_overlap: int) -> LoaderResult:
    from docx import Document

    document = Document(io.BytesIO(data))
    blocks: list[str] = [paragraph.text for paragraph in document.paragraphs]

    for table_number, table in enumerate(document.tables, start=1):
        rows = [[_format_cell(cell.text) for cell in row.cells] for row in table.rows]
        for _, section_text in _rows_to_sections(f"Table {table_number}", rows):
            blocks.append(section_text)

    blocks = [block for block in blocks if block.strip()]
    sections: list[tuple[str, str]] = []

    for start in range(0, len(blocks), PARAGRAPHS_PER_SECTION):
        group = blocks[start : start + PARAGRAPHS_PER_SECTION]
        sections.append(
            (
                f"Paragraphs {start + 1}-{start + len(group)}",
                "\n\n".join(group),
            )
        )

    return _split_sections(sections, chunk_size, chunk_overlap), len(sections)


def load_xlsx(data: bytes, chunk_size: int, chunk_overlap: int) -> LoaderResult:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sections: list[tuple[str, str]] = []

    try:
        for sheet in workbook.worksheets:
            rows = [
                [_format_cell(value) for value in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            sections.extend(_rows_to_sections(f"Sheet '{sheet.title}'", rows))
    finally:
        workbook.close()

    return _split_sections(sections, chunk_size, chunk_overlap), len(sections)


def load_csv(data: bytes, chunk_size: int, chunk_overlap: int) -> LoaderResult:
    text = data.decode("utf-8-sig", errors="replace")

    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel

    rows = [
        [cell.strip() for cell in row] for row in csv.reader(io.StringIO(text), dialect)
    ]
    sections = _rows_to_sections("CSV", rows)
    return _split_sections(sections, chunk_size, chunk_overlap), len(sections)


def load_pptx(data: bytes, chunk_size: int, chunk_overlap: int) -> LoaderResult:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    sections: list[tuple[str, str]] = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []

        for shape in slide.shapes:
            if (
                getattr(shape, "has_text_frame", False)
                and shape.text_frame.text.strip()
            ):
                texts.append(shape.text_frame.text)

            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    texts.append(
                        " | ".join(_format_cell(cell.text) for cell in row.cells)
                    )

        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                texts.append(f"Speaker notes: {notes}")

        sections.append((f"Slide {slide_number}", "\n".join(texts)))

    return _split_sections(sections, chunk_size, chunk_overlap), len(sections)


def load_text(data: bytes, chunk_size: int, chunk_overlap: int) -> LoaderResult:
    text = data.decode("utf-8-sig", errors="replace")
    sections = [("Text", text)]
    return _split_sections(sections, chunk_size, chunk_overlap), 1


LOADERS: dict[str, Callable[[bytes, int, int], LoaderResult]] = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".xlsx": load_xlsx,
    ".xlsm": load_xlsx,
    ".csv": load_csv,
    ".pptx": load_pptx,
    ".txt": load_text,
    ".md": load_text,
}

SUPPORTED_EXTENSIONS = tuple(extension.lstrip(".") for extension in LOADERS)


def load_document(
    filename: str,
    data: bytes,
    chunk_size: int = 1000,
    chunk_overlap: int = 150,
) -> LoaderResult:
    """Pick a loader from the file extension and return searchable chunks."""
    if not data:
        raise ValueError("Document content cannot be empty.")

    extension = PurePath(filename or "").suffix.lower()
    loader = LOADERS.get(extension)

    if loader is None:
        supported = ", ".join(sorted(LOADERS))
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{extension or filename}'. Supported: {supported}."
        )

    return loader(data, chunk_size, chunk_overlap)


def describe_location(chunk: DocumentChunk) -> str:
    """Return a readable source label for a chunk or search result."""
    return chunk.get("location") or f"Page {chunk['page']}"
