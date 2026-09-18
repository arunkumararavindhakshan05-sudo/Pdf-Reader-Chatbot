import io
from unittest.mock import patch

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from src.loaders import (
    MAX_ROWS_PER_CHUNK,
    SUPPORTED_EXTENSIONS,
    UnsupportedFileTypeError,
    describe_location,
    load_document,
)


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("Quarterly report for the Coimbatore branch.")
    document.add_paragraph("Revenue grew by twelve percent.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Revenue"
    table.cell(1, 0).text = "South"
    table.cell(1, 1).text = "12000"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx(row_count: int = 3) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["Month", "Total"])
    for number in range(1, row_count + 1):
        sheet.append([f"Month {number}", number * 100])
    workbook.create_sheet("Empty")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_pptx() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Launch plan"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "Ship the Android app first."
    slide.notes_slide.notes_text_frame.text = "Mention the APK link."
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_docx_includes_paragraphs_and_tables() -> None:
    chunks, sections = load_document("report.docx", make_docx())

    text = "\n".join(chunk["text"] for chunk in chunks)
    assert sections == 1
    assert "Coimbatore branch" in text
    assert "Region | Revenue" in text
    assert "South | 12000" in text
    assert chunks[0]["location"] == "Paragraphs 1-3"


def test_xlsx_repeats_header_and_labels_rows() -> None:
    chunks, sections = load_document("sales.XLSX", make_xlsx())

    assert sections == 1
    assert chunks[0]["location"] == "Sheet 'Sales', rows 2-4"
    assert chunks[0]["text"].splitlines()[0] == "Month | Total"
    assert "Month 3 | 300" in chunks[0]["text"]


def test_xlsx_splits_large_sheets_into_row_blocks() -> None:
    chunks, sections = load_document(
        "big.xlsx",
        make_xlsx(MAX_ROWS_PER_CHUNK + 5),
        chunk_size=10_000,
        chunk_overlap=0,
    )

    assert sections == 2
    assert [chunk["location"] for chunk in chunks] == [
        f"Sheet 'Sales', rows 2-{MAX_ROWS_PER_CHUNK + 1}",
        f"Sheet 'Sales', rows {MAX_ROWS_PER_CHUNK + 2}-{MAX_ROWS_PER_CHUNK + 6}",
    ]
    assert all(chunk["text"].startswith("Month | Total") for chunk in chunks)


def test_csv_detects_semicolon_delimiter() -> None:
    data = b"name;city\nArun;Coimbatore\nPriya;Chennai\n"

    chunks, _ = load_document("people.csv", data)

    assert chunks[0]["location"] == "CSV, rows 2-3"
    assert "Arun | Coimbatore" in chunks[0]["text"]


def test_pptx_reads_text_and_speaker_notes() -> None:
    chunks, sections = load_document("deck.pptx", make_pptx())

    assert sections == 1
    assert chunks[0]["location"] == "Slide 1"
    assert "Ship the Android app first." in chunks[0]["text"]
    assert "Speaker notes: Mention the APK link." in chunks[0]["text"]


def test_text_and_markdown_files() -> None:
    chunks, sections = load_document("notes.md", b"# Title\n\nHello world")

    assert sections == 1
    assert chunks[0]["location"] == "Text"
    assert "Hello world" in chunks[0]["text"]


def test_pdf_uses_existing_extractor_and_adds_page_label() -> None:
    fake_chunks = [{"text": "Page text", "page": 2, "chunk": 1}]

    with patch("src.loaders.extract_chunks", return_value=(fake_chunks, 4)):
        chunks, pages = load_document("file.pdf", b"%PDF")

    assert pages == 4
    assert chunks[0]["location"] == "Page 2"


def test_unsupported_extension_is_rejected() -> None:
    with pytest.raises(UnsupportedFileTypeError, match="Unsupported file type '.exe'"):
        load_document("setup.exe", b"MZ")


def test_empty_content_is_rejected() -> None:
    with pytest.raises(ValueError, match="cannot be empty"):
        load_document("notes.txt", b"")


def test_invalid_chunk_settings_are_rejected() -> None:
    with pytest.raises(ValueError, match="smaller than chunk size"):
        load_document("notes.txt", b"hello", chunk_size=10, chunk_overlap=10)


def test_describe_location_falls_back_to_page() -> None:
    assert describe_location({"text": "x", "page": 7, "chunk": 1}) == "Page 7"
    assert (
        describe_location({"text": "x", "page": 1, "chunk": 1, "location": "Slide 1"})
        == "Slide 1"
    )


def test_supported_extensions_cover_office_formats() -> None:
    assert {"pdf", "docx", "xlsx", "csv", "pptx", "txt", "md"} <= set(
        SUPPORTED_EXTENSIONS
    )
